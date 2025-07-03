

import os
import sys
import tempfile
import pytest
from src.handlers.txt_handler import TxtHandler
from src.handlers.pdf_handler import PdfHandler
from src.handlers.docx_handler import DocxHandler

print('PYTHON EXECUTABLE:', sys.executable)

class DummyDocx:
    @staticmethod
    def save(path):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Hello DOCX World!")
        doc.save(path)

def test_txt_handler():
    handler = TxtHandler()
    with tempfile.NamedTemporaryFile('w+', suffix='.txt', delete=False) as f:
        f.write("Hello TXT World!")
        f.flush()
        path = f.name
    try:
        text = handler.extract_text(path)
        assert "Hello TXT World!" in text
    finally:
        os.remove(path)

def test_pdf_handler():
    handler = PdfHandler()
    try:
        from fpdf import FPDF
    except ImportError:
        pytest.skip("fpdf not installed")
    with tempfile.NamedTemporaryFile('wb', suffix='.pdf', delete=False) as f:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, "Hello PDF World!", ln=True)
        pdf.output(f.name)
        path = f.name
    try:
        text = handler.extract_text(path)
        assert "Hello PDF World!" in text
    finally:
        os.remove(path)

def test_docx_handler():
    handler = DocxHandler()
    try:
        import docx
    except ImportError:
        pytest.skip("python-docx not installed")
    with tempfile.NamedTemporaryFile('wb', suffix='.docx', delete=False) as f:
        DummyDocx.save(f.name)
        path = f.name
    try:
        text = handler.extract_text(path)
        assert "Hello DOCX World!" in text
    finally:
        os.remove(path)
